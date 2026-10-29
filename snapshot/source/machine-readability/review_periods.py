"""Read exact originals to confirm population observation dates, not timestamps.

This does not change originals or mechanical scores. A recognized date header
must exist exactly once; every data row with a nonblank date must agree.
Missing dates are counted, and only a unique actual date is recorded.
"""
from __future__ import annotations
import argparse
import csv
import datetime as dt
import io
from pathlib import Path
import re
import unicodedata
from pipeline import ROOT,raw_bytes,read_json,write_json,now


def date_text(value):
    if isinstance(value,(dt.date,dt.datetime)):return value.date().isoformat() if isinstance(value,dt.datetime) else value.isoformat()
    s=unicodedata.normalize('NFKC',str(value)).strip()
    m=re.fullmatch(r'(\d{4})[-/](\d{1,2})[-/](\d{1,2})(?: 0:00:00)?',s)
    if m:
        try:return dt.date(*(int(x) for x in m.groups())).isoformat()
        except ValueError:return None
    return None


def confirm_population(data,fmt):
    if fmt in {'csv','tsv'}:
        for enc in ['utf-8-sig','cp932','utf-16']:
            try:text=data.decode(enc);break
            except UnicodeError:pass
        else:return {'status':'encoding_unconfirmed'}
        rows=list(csv.reader(io.StringIO(text),delimiter='\t' if fmt=='tsv' else ','))
    elif fmt in {'xlsx','xlsm'}:
        from openpyxl import load_workbook
        w=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        if len(w.worksheets)!=1:
            sheets=list(w.sheetnames);w.close()
            return {'status':'multiple_sheets_require_date_review','sheets':sheets}
        rows=list(w.worksheets[0].values);w.close()
    else:return {'status':'native_format_not_inspected'}
    if not rows:return {'status':'empty'}
    headers=[unicodedata.normalize('NFKC',str(v or '')).strip() for v in rows[0]]
    if headers.count('調査年月日')!=1:return {'status':'date_header_requires_review'}
    col=headers.index('調査年月日');dates=set();missing=invalid=0;data_rows=0
    for row in rows[1:]:
        if not any(v is not None and str(v).strip() for v in row):continue
        data_rows+=1;value=row[col] if col<len(row) else None
        if value is None or not str(value).strip():missing+=1;continue
        normalized=date_text(value)
        if normalized:dates.add(normalized)
        else:invalid+=1
    result={'status':'unique_observation_date_verified' if len(dates)==1 and not invalid else 'date_requires_review',
            'dates':sorted(dates),'date_header':'調査年月日','data_rows':data_rows,'rows_missing_date':missing,'rows_with_invalid_date':invalid,
            'method':'Read every populated record in first sheet; no catalog modification dates used'}
    if result['status']=='unique_observation_date_verified':result['observation_date']=next(iter(dates))
    return result


def finance_year_evidence(work, sources):
    """The national template's B1 title is direct evidence of the fiscal year."""
    from openpyxl import load_workbook
    evidence=[]
    for source in sources:
        file=work/'fetch'/(source['id']+'.json')
        if source['category']!='finance' or source['format'] not in {'xlsx','xlsm'} or not file.exists():continue
        fetched=read_json(file)
        if fetched['status']!='downloaded':continue
        w=load_workbook(io.BytesIO(raw_bytes(work,fetched)),read_only=True,data_only=True)
        sheet=w.worksheets[0];text=str(sheet['B1'].value or '')
        evidence.append({'source_id':source['id'],'original_sha256':fetched['sha256'],'sheet':sheet.title,'cell':'B1',
            'text':text,'target_fiscal_year':'FY2024','target_year_verified':'令和6年度' in unicodedata.normalize('NFKC',text)})
        w.close()
    return evidence


def main():
    p=argparse.ArgumentParser();p.add_argument('--work',type=Path,required=True);a=p.parse_args()
    sources=read_json(ROOT/'sources.json.gz');reviews=[]
    for s in sources:
        file=a.work/'fetch'/(s['id']+'.json')
        if s['category']!='population' or not file.exists():continue
        f=read_json(file)
        if f['status']!='downloaded':continue
        reviews.append({'source_id':s['id'],'original_sha256':f['sha256'],'checked_at':now(),**confirm_population(raw_bytes(a.work,f),s['format'])})
    write_json(a.work/'period-reviews.json',reviews)
    print('Population date checks:',len(reviews),'exact target date:',sum(r.get('observation_date')=='2026-01-01' and r.get('rows_missing_date')==0 for r in reviews))

if __name__=='__main__':main()
