#!/usr/bin/env python3
"""Normalize current facility-standard workbooks using the common MHLW schema."""
from __future__ import annotations

import argparse
import gzip
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

FIELDS = [
    "item_no", "prefecture_code", "prefecture", "category",
    "institution_number", "co_located_number", "symbol_number",
    "institution_name", "postal_code", "address", "phone", "fax",
    "licensed_beds_text", "accepted_name", "acceptance_symbol",
    "acceptance_number", "billing_start_date", "individual_start_date",
    "remark_heading", "remark_value", "municipality_code", "municipality",
]
CATEGORY_SLUG = {"医科": "medical", "歯科": "dental", "薬局": "pharmacy"}


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u3000", " ")).strip()


def workbook_records(content):
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = [clean(row[i]) if i < len(row) else "" for i in range(len(FIELDS))]
            if values[:2] == ["項番", "都道府県コード"]:
                continue
            if not values[1].isdigit() or values[3] not in CATEGORY_SLUG:
                continue
            if not values[4] or not values[7] or not values[13]:
                continue
            yield dict(zip(FIELDS, values))


def input_workbooks(input_root):
    for path in sorted(input_root.iterdir()):
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            yield path.name, path.read_bytes()
        elif suffix == ".zip":
            archive_bytes = path.read_bytes()
            with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
                for member in sorted(archive.namelist()):
                    if member.lower().endswith(".xlsx"):
                        yield f"{path.name}:{member}", archive.read(member)


def workbook_snapshot(content):
    found = set()
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=1, max_row=5, values_only=True):
            text = unicodedata.normalize("NFKC", " ".join(clean(value) for value in row if value is not None))
            match = re.search(r"令和\s*(\d+)\s*年\s*(\d+)\s*月\s*1\s*日\s*現在", text)
            if match:
                found.add(f"{2018 + int(match.group(1)):04d}-{int(match.group(2)):02d}")
    if len(found) > 1:
        raise RuntimeError(f"workbook contains multiple snapshot months: {sorted(found)}")
    return next(iter(found), "")


def snapshot_from_workbooks(input_root):
    found = {workbook_snapshot(content) for _, content in input_workbooks(input_root)}
    found.discard("")
    if not found:
        raise RuntimeError("no snapshot month found in workbook headers")
    return max(found)


def normalize(input_root, snapshot):
    grouped = {}
    sources = []
    for path in sorted(input_root.iterdir()):
        if path.suffix.lower() not in {".xlsx", ".zip"}:
            continue
        content = path.read_bytes()
        sources.append({
            "name": path.name,
            "sha256": hashlib.sha256(content).hexdigest(),
            "bytes": len(content),
        })
    for member_name, content in input_workbooks(input_root):
        if workbook_snapshot(content) != snapshot:
            continue
        for row in workbook_records(content):
            key = (
                row["prefecture_code"], row["category"], row["institution_number"],
                row["accepted_name"], row["acceptance_symbol"], row["acceptance_number"],
            )
            if key not in grouped:
                grouped[key] = {
                    field: row[field] for field in FIELDS
                    if field not in {"item_no", "remark_heading", "remark_value"}
                }
                grouped[key]["remarks"] = []
                grouped[key]["source_members"] = []
            if member_name not in grouped[key]["source_members"]:
                grouped[key]["source_members"].append(member_name)
            remark = "：".join(x for x in (row["remark_heading"], row["remark_value"]) if x)
            if remark and remark not in grouped[key]["remarks"]:
                grouped[key]["remarks"].append(remark)
    records = sorted(grouped.values(), key=lambda x: (
        x["prefecture_code"], x["category"], x["institution_number"],
        x["accepted_name"], x["acceptance_number"],
    ))
    used_sources = {
        member.split(":", 1)[0]
        for row in records for member in row["source_members"]
    }
    for row in records:
        del row["source_members"]
    sources = [source for source in sources if source["name"] in used_sources]
    return records, sources


def write_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")


def write_json_gzip(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    content = (json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n").encode()
    path.write_bytes(gzip.compress(content, compresslevel=9, mtime=0))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-root", required=True)
    parser.add_argument("--bureau", required=True)
    parser.add_argument("--snapshot")
    parser.add_argument("--output-root")
    args = parser.parse_args()
    input_root = Path(args.input_root)
    output_root = Path(args.output_root or f"data/facility-standards/normalized/{args.bureau}")
    snapshot = args.snapshot or snapshot_from_workbooks(input_root)
    records, sources = normalize(input_root, snapshot)
    if not records:
        raise RuntimeError("no records matched the common facility-standard schema")
    partitions = defaultdict(list)
    for row in records:
        partitions[(row["prefecture_code"], CATEGORY_SLUG[row["category"]])].append(row)
    files = []
    for (prefecture_code, category), subset in sorted(partitions.items()):
        path = output_root / snapshot / f"{prefecture_code}-{category}.json.gz"
        write_json_gzip(path, {
            "schemaVersion": 1, "bureau": args.bureau, "snapshot": snapshot,
            "prefectureCode": prefecture_code, "prefecture": subset[0]["prefecture"],
            "category": subset[0]["category"], "recordCount": len(subset), "records": subset,
        })
        files.append({
            "prefectureCode": prefecture_code, "prefecture": subset[0]["prefecture"],
            "category": subset[0]["category"], "recordCount": len(subset),
            "path": str(path.relative_to(output_root)),
        })
    index = {
        "schemaVersion": 1, "bureau": args.bureau, "snapshot": snapshot,
        "generatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "recordCount": len(records), "sourceFiles": sources, "partitions": files,
        "warning": "届出受理状況であり、診療実績、稼働状況、医療の質を示さない。",
    }
    write_json(output_root / "index.json", index)
    print(json.dumps({"recordCount": len(records), "partitions": len(files)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
